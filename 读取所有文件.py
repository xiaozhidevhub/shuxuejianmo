# -*- coding: utf-8 -*-
"""
读取所有文件.py
作用：遍历“D题”文件夹，把里面的 Word、Excel、CSV、MAT、TIF 文件内容
      读出来，统一写入“文件内容汇总.txt”，方便查看。
适合初学者：每一步都有注释。
"""

# ========== 第1步：导入需要用到的工具库 ==========
import os                  # os：操作文件和文件夹（比如遍历目录）
import pandas as pd        # pandas：读取表格数据（Excel、CSV）
import docx                # python-docx：读取 Word 文档（.docx）
import scipy.io as sio     # scipy.io：读取 MATLAB 数据文件（.mat）
import numpy as np         # numpy：做数值计算（统计最大值、最小值等）
from PIL import Image      # PIL：读取图片（这里用来读 .tif 高程图）

# ========== 第2步：设置路径 ==========
# 获取本脚本所在的文件夹，作为要遍历的根目录
根目录 = os.path.dirname(os.path.abspath(__file__))
# 输出结果保存到这个文本文件里
输出文件路径 = os.path.join(根目录, "文件内容汇总.txt")

# 让 pandas 打印表格时显示所有列、不自动换行
pd.set_option("display.max_columns", None)
pd.set_option("display.width", 250)


# ========== 第3步：为每种文件写一个读取函数 ==========

def 读取word(路径):
    """读取 .docx 文件：先读所有段落，再读所有表格"""
    文档 = docx.Document(路径)
    结果 = []
    # 逐段读取正文文字
    for 段落 in 文档.paragraphs:
        文字 = 段落.text.strip()      # strip() 去掉首尾空白
        if 文字 != "":                 # 空段落就跳过
            结果.append(文字)
    # 逐个读取 Word 中的表格
    for 序号, 表格 in enumerate(文档.tables, start=1):
        结果.append(f"\n[Word 表格 {序号}]")
        for 行 in 表格.rows:
            单元格文字 = [格.text.strip() for 格 in 行.cells]
            结果.append(" | ".join(单元格文字))
    return "\n".join(结果)


def 读取excel(路径):
    """读取 .xlsx 文件的每一个工作表，全部打印出来"""
    结果 = []
    # sheet_name=None 表示读取所有工作表，返回一个字典 {表名: 数据}
    所有表 = pd.read_excel(路径, sheet_name=None, header=None)
    for 表名, 数据 in 所有表.items():
        结果.append(f"[工作表：{表名}]  共 {数据.shape[0]} 行 {数据.shape[1]} 列")
        结果.append(数据.to_string())
    return "\n".join(结果)


def 读取csv(路径):
    """读取 .csv 文件：数据可能很大，只显示基本信息和前 10 行"""
    # 中文 CSV 常见编码有 utf-8 和 gbk，挨个尝试
    for 编码 in ["utf-8-sig", "gbk"]:
        try:
            数据 = pd.read_csv(路径, encoding=编码)
            break                      # 读取成功就跳出循环
        except UnicodeDecodeError:
            continue                   # 编码不对就试下一个
    结果 = [f"共 {数据.shape[0]} 行 {数据.shape[1]} 列",
          "列名：" + ", ".join(str(列) for 列 in 数据.columns),
          "前 10 行：",
          数据.head(10).to_string()]
    return "\n".join(结果)


def 读取mat(路径):
    """读取 .mat 文件：列出里面有哪些变量，以及每个变量的形状"""
    数据 = sio.loadmat(路径)
    结果 = []
    for 变量名, 值 in 数据.items():
        if 变量名.startswith("__"):   # "__header__" 等是系统信息，跳过
            continue
        结果.append(f"变量 {变量名}：类型 {type(值).__name__}，形状 {getattr(值, 'shape', '')}")
    return "\n".join(结果)


def 读取tif(路径):
    """读取 .tif 高程图：显示尺寸和高程的最小/最大/平均值"""
    图片 = Image.open(路径)
    矩阵 = np.array(图片, dtype=float)  # 转成数字矩阵，每个格子是一个高程值
    return (f"尺寸（行×列）：{矩阵.shape}\n"
            f"最小值：{np.nanmin(矩阵):.2f}  最大值：{np.nanmax(矩阵):.2f}  "
            f"平均值：{np.nanmean(矩阵):.2f}")


# 用字典把“文件后缀”和“对应的读取函数”对应起来
读取函数表 = {
    ".docx": 读取word,
    ".xlsx": 读取excel,
    ".csv": 读取csv,
    ".mat": 读取mat,
    ".tif": 读取tif,
}

# ========== 第4步：遍历文件夹，逐个读取并写入结果 ==========
with open(输出文件路径, "w", encoding="utf-8") as 输出:
    # os.walk 会自动进入每一层子文件夹
    for 当前目录, 子目录列表, 文件列表 in os.walk(根目录):
        for 文件名 in 文件列表:
            后缀 = os.path.splitext(文件名)[1].lower()   # 取出文件后缀，如 ".xlsx"
            if 后缀 not in 读取函数表:                   # 不认识的文件类型跳过
                continue
            完整路径 = os.path.join(当前目录, 文件名)
            相对路径 = os.path.relpath(完整路径, 根目录)
            输出.write("\n" + "=" * 80 + "\n")
            输出.write(f"文件：{相对路径}\n")
            输出.write("=" * 80 + "\n")
            try:
                内容 = 读取函数表[后缀](完整路径)          # 调用对应的读取函数
                输出.write(内容 + "\n")
            except Exception as 错误:                     # 出错也不中断，记录下来
                输出.write(f"读取失败：{错误}\n")

print("读取完成，结果已保存到：", 输出文件路径)
